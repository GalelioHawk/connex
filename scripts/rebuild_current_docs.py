from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
GREEN = RGBColor(0x00, 0x86, 0x5A)
INK = RGBColor(0x13, 0x1A, 0x17)
MUTED = RGBColor(0x63, 0x6B, 0x67)
PALE = "EAF6F0"


def shade(cell, fill):
    props = cell._tc.get_or_add_tcPr()
    element = OxmlElement("w:shd")
    element.set(qn("w:fill"), fill)
    props.append(element)


def set_cell_margins(cell):
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for edge, value in (("top", 80), ("start", 120), ("bottom", 80), ("end", 120)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def configure(doc, title):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    for name, size, before, after in (("Heading 1", 16, 16, 8), ("Heading 2", 13, 12, 6), ("Heading 3", 12, 8, 4)):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = GREEN
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
    header = section.header.paragraphs[0]
    header.text = title.upper()
    header.style = doc.styles["Normal"]
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = MUTED
    footer = section.footer.paragraphs[0]
    footer.text = "Connex Digital (Pty) Ltd | Internal working document | July 2026"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.runs[0].font.size = Pt(8)
    footer.runs[0].font.color.rgb = MUTED


def title_block(doc, kicker, title, subtitle):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(kicker.upper())
    r.bold = True; r.font.size = Pt(10); r.font.color.rgb = GREEN
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(title)
    r.bold = True; r.font.size = Pt(27); r.font.color.rgb = INK
    p = doc.add_paragraph(subtitle)
    p.paragraph_format.space_after = Pt(18)
    p.runs[0].font.size = Pt(13); p.runs[0].font.color.rgb = MUTED


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for idx, header in enumerate(headers):
        cell = t.rows[0].cells[idx]
        cell.text = header
        shade(cell, PALE)
        cell.paragraphs[0].runs[0].bold = True
    for row in rows:
        cells = t.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = str(value)
    for row in t.rows:
        for idx, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if widths:
                cell.width = Inches(widths[idx])
    return t


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def build_technical():
    doc = Document()
    configure(doc, "Connex Technical Development Plan")
    title_block(doc, "Current architecture", "Connex Technical Development Plan", "Expo + Convex build guide and alpha-readiness reference")
    table(doc, ["Field", "Current value"], [
        ["Status date", "12 July 2026"], ["Active branch", "dev"],
        ["Mobile", "Expo 55 / React Native 0.83 / TypeScript"],
        ["Backend", "Convex database, realtime, functions, storage, cron"],
        ["Authentication", "Phone/password, bcrypt actions, 30-day Convex sessions"],
        ["Distribution", "Expo EAS; local physical-iPhone development build validated"],
    ], [1.8, 4.7])
    doc.add_heading("1. System of record", level=1)
    doc.add_paragraph("The active product is apps/mobile. All server logic lives in apps/mobile/convex. The former Express, Supabase, Railway, Redis, REST, JWT, Axios, and NativeWind architecture has been removed and must not be recreated.")
    doc.add_heading("2. Architecture", level=1)
    table(doc, ["Layer", "Implementation", "Responsibility"], [
        ["Client", "Expo React Native", "iOS/Android screens, navigation, device APIs"],
        ["Reactive data", "Convex useQuery", "Messages, conversations, Feed, Edu, status, calls"],
        ["Writes", "Convex mutations", "Validated transactional state changes"],
        ["Side effects", "Convex actions", "bcrypt, Agora tokens, external APIs, push delivery"],
        ["Local state", "Zustand", "Authentication, preferences, current Chat tab"],
        ["Files", "Convex storage", "Avatars, chat media, status media, Edu PDFs"],
    ], [1.1, 1.7, 3.7])
    doc.add_heading("3. Implemented Phase 1 surfaces", level=1)
    bullets(doc, [
        "Chat: direct and group conversations, message requests, media, delivery/read state, presence, deletion, and push scheduling.",
        "Group administration: rename, add/remove members, leave group, and replacement-admin promotion.",
        "Calls: Agora voice/video signalling, incoming call handling, call history, and missed-call state.",
        "Status: 24-hour text/media updates, views, deletion, and authenticated storage uploads.",
        "SOS: mutual-contact requests, three-second hold interaction, emergency system messages, push scheduling, and activation limits.",
        "Feed: community/alerts tabs, area/category filtering, posting, likes, and official loadshedding summary.",
        "Edu: papers, PDF viewer, notes, tutors, contact requests, questions, bookmarks/recent-view backend, and idempotent starter data.",
    ])
    doc.add_heading("4. Security model", level=1)
    bullets(doc, [
        "Every protected function accepts a Convex session ID and validates expiry and active-user status.",
        "Admin catalogue mutations use role checks; production seeding is admin-only.",
        "Storage upload URLs require authentication; conversation media additionally requires membership.",
        "Rate limits cover authentication, messaging, Feed posting, Edu activity, tutor requests, and SOS activation.",
        "Expo push tokens are stored as expoPushToken; fcmToken remains read-only for migration compatibility.",
    ])
    doc.add_heading("5. Required environment configuration", level=1)
    table(doc, ["Location", "Variable", "Purpose"], [
        ["Mobile", "EXPO_PUBLIC_CONVEX_URL", "Convex deployment URL"],
        ["Mobile", "EXPO_PUBLIC_APP_ENV", "Client environment label"],
        ["Mobile", "EXPO_PUBLIC_AGORA_APP_ID", "Agora client initialization"],
        ["Convex", "APP_ENV", "Production gates"],
        ["Convex", "ESKOMSEPUSH_API_KEY", "Official loadshedding refresh"],
        ["Convex", "AGORA_APP_ID / AGORA_APP_CERTIFICATE", "Server RTC token generation"],
    ], [1.0, 2.2, 3.3])
    doc.add_heading("6. Release checks", level=1)
    bullets(doc, [
        "Run npm run release:check from apps/mobile.",
        "Verify register/login/logout and forced session expiry.",
        "Test message, call, and SOS push delivery in foreground, background, and terminated states on two physical devices.",
        "Test light/dark mode and permission prompts on iOS and Android.",
        "Deploy Convex schema/functions to a non-production deployment before production promotion.",
        "Complete moderation, privacy, legal, store metadata, screenshots, and support workflows before public submission.",
    ])
    doc.add_heading("7. Phase boundaries", level=1)
    doc.add_paragraph("Phase 2 Pay, Phase 3 Clips, and Phase 4 Logistics remain roadmap concepts. Do not add payment, wallet, financial, mapping, ride, delivery, or video infrastructure to Phase 1 without a separate approved architecture and compliance plan.")
    out = ROOT / "docs/technical/Connex_Technical_Development_Plan.docx"
    doc.save(out)


def build_business(path, report=False):
    doc = Document()
    configure(doc, "Connex Business Plan" if not report else "Connex Business Report")
    title_block(doc, "Confidential working plan", "Connex Business Plan" if not report else "Connex Business & Readiness Report", "Africa-first communication, community safety, and education platform")
    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph("Connex Digital (Pty) Ltd is building a mobile platform that combines real-time communication, community alerts, emergency-contact SOS, and free learner resources. Phase 1 is in alpha hardening: the core Expo/Convex product is implemented, while multi-device service validation, moderation, legal review, and public-store readiness remain outstanding.")
    table(doc, ["Product", "Phase 1 value", "Current state"], [
        ["Chat", "Everyday communication and groups", "Implemented; alpha QA required"],
        ["Feed & Alerts", "Local community information", "Implemented; content/moderation depth required"],
        ["Edu", "Grades 10-12 papers, notes, tutors, help", "Implemented; verified paper URLs need expansion"],
        ["SOS", "Trusted-contact emergency notification", "Implemented; delivery must be tested on multiple devices"],
    ], [1.2, 2.8, 2.5])
    doc.add_heading("Problem and positioning", level=1)
    doc.add_paragraph("Connex addresses fragmentation across messaging, local information, safety coordination, and learner resources. The launch strategy should prove one focused community use case at a time rather than claim continental scale before retention and trust are demonstrated.")
    doc.add_heading("Go-to-market hypothesis", level=1)
    bullets(doc, [
        "Begin with a controlled Heidelberg/Gauteng alpha using known testers and community partners.",
        "Use Edu resources and practical community alerts as utility-led acquisition channels.",
        "Measure activation, seven-day retention, successful conversations, useful alerts, paper opens, and SOS delivery reliability.",
        "Expand only after privacy, moderation, support, abuse response, and service reliability meet documented thresholds.",
    ])
    doc.add_heading("Business model", level=1)
    doc.add_paragraph("Phase 1 remains free and focused on trust and retention. Pay, Clips, advertising, tutoring monetisation, and Logistics are planning options rather than committed revenue. Each requires separate validation; Pay additionally requires licensed partners and regulatory advice before handling money.")
    doc.add_heading("Current operating priorities", level=1)
    table(doc, ["Priority", "Outcome", "Gate"], [
        ["Alpha reliability", "Two-device end-to-end QA", "No critical auth/chat/SOS failures"],
        ["Safety", "Moderation and incident process", "Named owner and response playbook"],
        ["Edu catalogue", "Verified official paper links", "No fabricated or broken sources"],
        ["Legal", "POPIA/privacy/terms review", "Counsel-reviewed public documents"],
        ["Growth", "Small measured pilot", "Retention evidence before paid acquisition"],
    ], [1.2, 2.8, 2.5])
    doc.add_heading("Funding approach", level=1)
    doc.add_paragraph("The historic R500,000-R2,000,000 range remains a scenario, not a valuation or committed round. Before outreach, build a bottom-up 12-18 month budget tied to provider costs, devices, legal/compliance, moderation/support, and a narrowly defined pilot. Investor materials must distinguish implemented code, device-tested behavior, and future roadmap concepts.")
    doc.add_heading("Risks", level=1)
    bullets(doc, [
        "Safety-critical SOS behavior creates a higher duty of care and must never be marketed as emergency-service dispatch.",
        "A broad super-app story can dilute focus; Phase 1 needs a measurable initial wedge.",
        "Community content requires reporting, moderation, appeals, and abuse controls.",
        "Education sources must remain verified, lawful, and maintainable.",
        "Financial and market projections in earlier documents are stale and must not be presented as audited forecasts.",
    ])
    doc.add_heading("Roadmap", level=1)
    table(doc, ["Stage", "Objective"], [
        ["Alpha hardening", "Security, service credentials, moderation, legal review, two-device QA"],
        ["Closed pilot", "Small invited community, analytics, support loop, retention measurement"],
        ["Public Phase 1", "Store release only after reliability and trust gates pass"],
        ["Later phases", "Separate discovery, business case, architecture, and compliance approval"],
    ], [1.5, 5.0])
    doc.add_paragraph("Planning note: market sizes, competitor metrics, projections, deadlines, and funding programme details require current primary-source verification before external use.")
    doc.save(path)


if __name__ == "__main__":
    build_technical()
    build_business(ROOT / "docs/business/Connex_Business_Plan.docx")
    build_business(ROOT / "docs/business/Connex_Business_Report.docx", report=True)
